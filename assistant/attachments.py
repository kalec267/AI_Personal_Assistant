"""Read selected local documents as data; never execute embedded content."""
import base64
from dataclasses import dataclass
import io
import json
import os
from pathlib import Path
import subprocess
import warnings
import zipfile

from .core import UserError

MAX_FILES = 5
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT = 40000
MAX_TOTAL_TEXT = 60000
MAX_IMAGES = 4
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
DOCUMENT_TYPES = {".pdf", ".pptx", ".ppt", ".xlsx", ".xls", ".docx", ".txt", ".md", ".csv", ".json"}


@dataclass(frozen=True)
class Attachment:
    name: str
    text: str = ""
    image: str = ""
    note: str = ""


class TextCollector:
    def __init__(self, limit=MAX_TEXT):
        self.parts = []
        self.length = 0
        self.limit = limit
        self.truncated = False

    def add(self, text):
        text = str(text).strip()
        if not text:
            return
        remaining = self.limit - self.length
        if len(text) + 1 > remaining:
            self.truncated = True
        value = (text + "\n")[:remaining]
        self.parts.append(value)
        self.length += len(value)

    @property
    def full(self):
        return self.length >= self.limit

    def result(self):
        return "".join(self.parts).strip()


def check_zip(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 120 * 1024 * 1024:
            raise UserError("문서의 압축 해제 크기가 너무 큽니다. 필요한 부분만 별도 파일로 저장하세요.")


def pdf_text(path, out):
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise UserError("암호화된 PDF입니다. 암호를 해제한 사본을 첨부하세요.")
    for index, page in enumerate(reader.pages):
        if index >= 200 or out.full:
            out.truncated = True
            break
        text = page.extract_text() or ""
        if text.strip():
            out.add(f"[페이지 {index + 1}]\n{text}")
    if not out.result():
        raise UserError("읽을 수 있는 PDF 텍스트가 없어요. 스캔 문서는 페이지를 이미지로 첨부하거나 OCR 처리한 PDF를 사용하세요.")


def pptx_text(path, out):
    from pptx import Presentation
    def shapes(items):
        for shape in items:
            if out.full:
                return
            if shape.has_text_frame:
                out.add(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    out.add(" | ".join(cell.text for cell in row.cells))
                    if out.full:
                        break
            if getattr(shape, "has_chart", False):
                for series in shape.chart.series:
                    out.add(f"[차트] {series.name}: {list(series.values)}")
            if hasattr(shape, "shapes"):
                shapes(shape.shapes)
    presentation = Presentation(path)
    for index, slide in enumerate(presentation.slides):
        if index >= 200 or out.full:
            out.truncated = True
            break
        out.add(f"[슬라이드 {index + 1}]")
        shapes(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes and notes.text.strip():
                out.add("[발표자 메모] " + notes.text)


def legacy_ppt_text(path, out):
    if os.name != "nt":
        raise UserError(".ppt는 Windows의 설치된 Microsoft PowerPoint가 필요합니다. .pptx로 저장하면 Office 없이 읽을 수 있어요.")
    script = Path(__file__).with_name("document.ps1")
    try:
        result = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-File", str(script)],
                                input=json.dumps({"path": str(path.resolve())}), capture_output=True, encoding="utf-8",
                                timeout=45, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), check=True)
        value = json.loads(result.stdout)
        if not isinstance(value, dict) or not isinstance(value.get("text"), str):
            raise ValueError()
        out.add(value["text"])
        out.truncated |= bool(value.get("truncated"))
    except (OSError, ValueError, subprocess.SubprocessError):
        raise UserError(".ppt를 읽지 못했어요. Microsoft PowerPoint 설치와 실행 정책을 확인하거나 .pptx로 저장해 첨부하세요.") from None


def excel_text(path, out):
    from openpyxl import load_workbook
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        for sheet in formulas.worksheets[:10]:
            out.add(f"[시트 {sheet.title}]")
            if sheet.max_row is None or sheet.max_column is None or sheet.max_row > 1000 or sheet.max_column > 40:
                out.truncated = True
            raw_rows = sheet.iter_rows(max_row=min(sheet.max_row or 1000, 1000), max_col=min(sheet.max_column or 40, 40), values_only=True)
            cached_rows = cached[sheet.title].iter_rows(max_row=min(sheet.max_row or 1000, 1000), max_col=min(sheet.max_column or 40, 40), values_only=True)
            for index, (raw, values) in enumerate(zip(raw_rows, cached_rows), 1):
                if out.full:
                    out.truncated = True
                    return
                cells = [str(v if v is not None else f"[계산값 없는 수식: {r}]" if isinstance(r, str) and r.startswith("=") else r or "") for r, v in zip(raw, values)]
                if any(cells):
                    out.add(f"행 {index}: " + " | ".join(cells))
        if len(formulas.worksheets) > 10:
            out.truncated = True
    finally:
        formulas.close()
        if cached:
            cached.close()


def xls_text(path, out):
    import xlrd
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        for sheet in book.sheets()[:10]:
            out.add(f"[시트 {sheet.name}]")
            if sheet.nrows > 1000 or sheet.ncols > 40:
                out.truncated = True
            for index in range(min(sheet.nrows, 1000)):
                if out.full:
                    out.truncated = True
                    return
                out.add(f"행 {index + 1}: " + " | ".join(str(value) for value in sheet.row_values(index, end_colx=min(sheet.ncols, 40))))
        out.truncated |= book.nsheets > 10
    finally:
        book.release_resources()


def docx_text(path, out):
    from docx import Document
    from docx.table import Table
    document = Document(path)
    for block in document.iter_inner_content():
        if out.full:
            out.truncated = True
            break
        if isinstance(block, Table):
            out.add("[표]")
            for row in block.rows:
                out.add(" | ".join(cell.text for cell in row.cells))
                if out.full:
                    break
        else:
            out.add(block.text)


def image_attachment(path):
    from PIL import Image, ImageOps
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path) as original:
            if original.width * original.height > 40_000_000:
                raise UserError("이미지는 4천만 화소 이하로 줄여 첨부하세요.")
            image = ImageOps.exif_transpose(original)
            notes = []
            if getattr(original, "n_frames", 1) > 1:
                notes.append("여러 프레임 중 첫 이미지만 분석합니다.")
            if max(image.size) > 1600:
                notes.append("긴 변 1600픽셀로 축소했습니다. 작은 글자는 확대해서 다시 첨부하세요.")
            image.thumbnail((1600, 1600))
            rgba = image.convert("RGBA")
            flat = Image.new("RGB", rgba.size, "white")
            flat.paste(rgba, mask=rgba.getchannel("A"))
            output = io.BytesIO()
            flat.save(output, format="JPEG", quality=90)
    return Attachment(path.name, image=base64.b64encode(output.getvalue()).decode("ascii"), note=" ".join(notes))


def read_file(filename):
    path = Path(filename)
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_FILE_BYTES:
        raise UserError("비어 있지 않은 20MB 이하의 파일을 선택하세요.")
    suffix = path.suffix.lower()
    if suffix not in DOCUMENT_TYPES | IMAGE_TYPES:
        raise UserError("지원하지 않는 형식입니다. .pptx, .xlsx, .docx, PDF, 텍스트 또는 이미지를 첨부하세요.")
    try:
        if suffix in IMAGE_TYPES:
            return image_attachment(path)
        if suffix in {".pptx", ".xlsx", ".docx"}:
            check_zip(path)
        out = TextCollector()
        readers = {".pdf": pdf_text, ".pptx": pptx_text, ".ppt": legacy_ppt_text, ".xlsx": excel_text, ".xls": xls_text, ".docx": docx_text}
        if suffix in readers:
            readers[suffix](path, out)
        else:
            raw = path.read_bytes()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = raw.decode("cp949")
                except UnicodeDecodeError:
                    raise UserError("텍스트 인코딩을 읽을 수 없어요. UTF-8로 저장해 첨부하세요.") from None
            out.add(text)
        if not out.result():
            raise UserError("분석할 텍스트를 찾지 못했어요. 이미지로 된 내용은 이미지 파일로 첨부하세요.")
        note = ""
        if suffix in {".pdf", ".ppt", ".pptx", ".docx"}:
            note = "추출한 텍스트·표를 분석합니다. 문서 안의 이미지·도형은 이미지로 별도 첨부하세요."
        if suffix in {".xlsx", ".xls"}:
            note = "수식은 실행하지 않으며 저장된 계산값을 사용합니다. 숨긴 시트도 포함됩니다."
        if out.truncated:
            note += " 읽기 제한으로 일부 내용만 포함했습니다 (문서 최대 40,000자, PDF/PPT 200페이지·슬라이드, Excel 10시트·1,000행·40열)."
        return Attachment(path.name, text=out.result(), note=note)
    except UserError:
        raise
    except ImportError:
        raise UserError("문서 처리 패키지가 필요합니다. 최신 setup.bat을 다시 실행하세요.") from None
    except Exception:
        raise UserError("파일을 읽지 못했어요. 손상·암호화 여부와 확장자를 확인하세요.") from None


def validate_attachments(items):
    if not isinstance(items, (list, tuple)) or not 1 <= len(items) <= MAX_FILES or any(not isinstance(item, Attachment) for item in items):
        raise UserError("첨부는 한 번에 1~5개까지 사용할 수 있어요.")
    if sum(len(item.text) for item in items) > MAX_TOTAL_TEXT:
        raise UserError("첨부 문서의 텍스트 합계가 60,000자를 넘습니다. 파일을 나누어 분석하세요.")
    if sum(bool(item.image) for item in items) > MAX_IMAGES:
        raise UserError("이미지는 한 번에 4개까지 분석할 수 있어요.")


def read_files(filenames):
    items, errors = [], []
    if len(filenames) > MAX_FILES:
        raise UserError("첨부 파일을 5개 이하로 선택하세요.")
    for filename in filenames:
        try:
            item = read_file(filename)
            validate_attachments(items + [item])
            items.append(item)
        except (UserError, OSError) as error:
            errors.append(f"{Path(filename).name}: {error}")
    return {"items": items, "errors": errors}
